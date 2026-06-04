#!/usr/bin/env python3
"""
SOA Stock — Excel to JSON Converter
วิธีใช้: python convert_xlsx.py <ชื่อไฟล์.xlsx>
"""
import sys
import json
import re
from datetime import datetime, date

try:
    import openpyxl
except ImportError:
    import subprocess
    subprocess.run([sys.executable, '-m', 'pip', 'install', 'openpyxl', '--break-system-packages', '-q'])
    import openpyxl

# Column mapping (Thai/English)
COL_MAP = {
    'regional': ['Regional', 'regional'],
    'store': ['Store', 'store'],
    'zone': ['Zone', 'zone'],
    'dept': ['Dept', 'dept'],
    'item': ['Item', 'item'],
    'desc': ['Description', 'description', 'ชื่อสินค้า'],
    'price': ['ราคาสินค้า (฿)', 'Price', 'price'],
    'diff': ['Diff', 'diff'],
    'value': ['มูลค่า Diff (฿)', 'Value', 'value'],
    'remark': ['Remark', 'remark'],
    'status': ['Status', 'status'],
    'date': ['วันที่บันทึก', 'Date', 'date'],
}

def fmt_date(v):
    if v is None:
        return ''
    if isinstance(v, (datetime, date)):
        y = v.year
        if y > 2500:
            y -= 543
        return f'{y}-{v.month:02d}-{v.day:02d}'
    s = str(v).strip()
    m = re.search(r'(\d{1,2})[\/\-\.](\d{1,2})[\/\-\.](\d{4})', s)
    if m:
        y = int(m.group(3))
        if y > 2500:
            y -= 543
        return f'{y}-{m.group(2).zfill(2)}-{m.group(1).zfill(2)}'
    return ''

def convert(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    headers = [str(c.value).strip() if c.value else '' for c in next(ws.iter_rows(min_row=1, max_row=1))]

    # Build column index mapping
    cm = {}
    for key, candidates in COL_MAP.items():
        for candidate in candidates:
            if candidate in headers:
                cm[key] = headers.index(candidate)
                break

    def get(row, key):
        idx = cm.get(key)
        if idx is None:
            return ''
        v = row[idx].value
        return v if v is not None else ''

    rows = []
    for row in ws.iter_rows(min_row=2):
        regional = str(get(row, 'regional')).strip()
        if not regional or regional.lower() == 'none':
            continue
        store = str(get(row, 'store')).strip().zfill(3)
        rows.append({
            'regional': regional,
            'store': store,
            'zone': str(get(row, 'zone')).strip(),
            'dept': str(get(row, 'dept')).strip(),
            'item': str(get(row, 'item')).strip(),
            'desc': str(get(row, 'desc')).strip(),
            'price': float(get(row, 'price') or 0),
            'diff': float(get(row, 'diff') or 0),
            'value': float(get(row, 'value') or 0),
            'remark': str(get(row, 'remark') or '').strip(),
            'status': str(get(row, 'status')).strip(),
            'dateStr': fmt_date(get(row, 'date')),
        })

    return rows

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('Usage: python convert_xlsx.py <file.xlsx>')
        sys.exit(1)
    
    xlsx_path = sys.argv[1]
    rows = convert(xlsx_path)
    
    out_path = 'data.json'
    with open(out_path, 'w', encoding='utf-8') as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)
    
    print(f'✓ แปลงสำเร็จ: {len(rows)} rows → {out_path}')
